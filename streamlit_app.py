import streamlit as st
import io, re
import pandas as pd
import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

st.set_page_config(page_title="IATA BSP Converter PRO | Mahmoud Amin", page_icon="✈️", layout="wide")

st.markdown("""
<style>
header {visibility: hidden;}
.stApp { background-color: #0F172A; }
.header-box { background: linear-gradient(135deg, #1E3A8A 0%, #3B82F6 100%); padding: 35px; border-radius: 20px; text-align: center; margin-bottom: 25px; }
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="header-box">
<h1 style="color:white; margin:0;">✈️ IATA / BSP CONVERTER</h1>
<p style="color:#BFDBFE;">FCAGBILLDET PDF → EXCEL • PROFESSIONAL EDITION v4 FINAL</p>
<p style="color:#93C5FD; font-style:italic;">Developed by Mahmoud Amin • 2026</p>
</div>
""", unsafe_allow_html=True)

def parse_pdf(pdf_bytes):
    tickets = []
    full_text = ""
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            txt = page.extract_text() or ""
            full_text += txt + "\n"
            # tables
            for tbl in (page.extract_tables() or []):
                for row in tbl:
                    if not row: continue
                    line = " ".join([str(c or "") for c in row])
                    if re.search(r'\d{3}[-\s]?\d{10,13}', line):
                        tickets.append({"doc_no": re.search(r'(\d{3}[-\s]?\d{10,13})', line).group(0), "raw": line[:250]})
    if not tickets:
        for line in full_text.split("\n"):
            if re.search(r'\d{3}[-\s]?\d{10,13}', line) and len(line.strip())>15:
                tickets.append({"doc_no": re.search(r'(\d{3}[-\s]?\d{10,13})', line).group(0), "raw": line.strip()[:250]})
    if not tickets and full_text:
        # fallback: show all lines
        for l in [x.strip() for x in full_text.split("\n") if len(x.strip())>20][:300]:
            tickets.append({"doc_no":"LINE", "raw": l[:250]})
    return tickets, full_text

def make_excel(tickets, full_text, filename):
    out = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "BSP-ISSUES"
    fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    ws.append(["#", "Document No", "Extracted Line"])
    for c in ws[1]: c.fill=fill; c.font=font
    for i,t in enumerate(tickets,1): ws.append([i, t["doc_no"], t["raw"]])
    ws.column_dimensions['A'].width=6
    ws.column_dimensions['B'].width=22
    ws.column_dimensions['C'].width=100
    ws2 = wb.create_sheet("FULL_TEXT")
    ws2.append([full_text[:32000]])
    ws3 = wb.create_sheet("INFO")
    ws3.append(["Developed by", "Mahmoud Amin"])
    ws3.append(["File", filename])
    ws3.append(["Records", len(tickets)])
    wb.save(out); out.seek(0)
    return out

c1,c2 = st.columns([2,1], gap="large")
with c1:
    up = st.file_uploader("Select PDF File", type=["pdf"])
    if up:
        st.success(f"Selected: {up.name} ({up.size//1024} KB)")
        with st.spinner("Converting..."):
            tickets, text = parse_pdf(up.getvalue())
        st.metric("Records Found", len(tickets))
        if tickets:
            st.dataframe(pd.DataFrame(tickets[:50]), use_container_width=True)
            xls = make_excel(tickets, text, up.name)
            st.download_button("⚡ GENERATE & DOWNLOAD EXCEL NOW", xls, file_name=up.name.replace(".pdf","_CONVERTED.xlsx"), mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            st.balloons()
            st.success(f"✅ Done! {len(tickets)} records - Developed by Mahmoud Amin")
with c2:
    st.info("**Fixed v4 FINAL**\n\nNo fitz dependency\n\nRequirements:\nstreamlit\npdfplumber\nopenpyxl\npandas\n\nPut 2 files in GitHub root:\n- streamlit_app.py\n- requirements.txt")
