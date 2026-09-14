"""
IATA BSP Converter PRO - Streamlit Edition
Developed by Mahmoud Amin
Run: streamlit run streamlit_app.py
Deploy to: streamlit.io Cloud for a public link
"""
import streamlit as st
import os, re, io
import pandas as pd

# Page Config - Fullscreen Professional
st.set_page_config(
    page_title="IATA BSP Converter PRO | Mahmoud Amin",
    page_icon="✈️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# Professional CSS - Dark Blue & Modern
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;700&family=Inter:wght@400;700;900&display=swap');
    .main { background-color: #0F172A; }
    header {visibility: hidden;}
    .stApp { background-color: #0F172A; }
    .header-box {
        background: linear-gradient(135deg, #1E3A8A 0%, #1E40AF 50%, #3B82F6 100%);
        padding: 40px 30px; border-radius: 20px; text-align: center;
        margin-bottom: 30px; box-shadow: 0 20px 40px rgba(59,130,246,0.3);
    }
    .header-title { font-family: 'Inter'; font-size: 42px; font-weight: 900; color: white; letter-spacing: -1px; }
    .header-sub { font-family: 'Inter'; font-size: 16px; color: #BFDBFE; margin-top: 8px; letter-spacing: 3px; }
    .header-dev { font-family: 'Cairo'; font-size: 14px; color: #93C5FD; margin-top: 10px; font-style: italic; }
    .card { background: #1E293B; padding: 25px; border-radius: 15px; border: 1px solid #334155; box-shadow: 0 10px 30px rgba(0,0,0,0.3); }
    .step-label { color: #3B82F6; font-weight: 800; font-size: 12px; letter-spacing: 2px; }
    .stButton>button { background: linear-gradient(135deg, #10B981, #059669); color: white; font-weight: 900; font-size: 18px; padding: 18px; border-radius: 12px; border: none; width: 100%; box-shadow: 0 10px 20px rgba(16,185,129,0.3); }
    .stButton>button:hover { background: linear-gradient(135deg, #059669, #047857); transform: translateY(-2px); }
    .info-box { background: #334155; padding: 20px; border-radius: 12px; color: #F1F5F9; font-family: monospace; }
    .footer { text-align: center; color: #64748B; padding: 30px; font-size: 13px; margin-top: 40px; }
</style>
""", unsafe_allow_html=True)

# Header
st.markdown("""
<div class="header-box">
    <div class="header-title">✈️ IATA / BSP CONVERTER</div>
    <div class="header-sub">FCAGBILLDET PDF → EXCEL • PROFESSIONAL EDITION</div>
    <div class="header-dev">Developed by Mahmoud Amin • 2026</div>
</div>
""", unsafe_allow_html=True)

# Try import parsers
try:
    import pdfplumber
    HAS_PDFPLUMBER = True
except:
    HAS_PDFPLUMBER = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_OPENPYXL = True
except:
    HAS_OPENPYXL = False

def parse_pdf_to_data(pdf_file):
    """Parse IATA PDF and extract tickets"""
    tickets = []
    summary_text = ""
    try:
        if HAS_PDFPLUMBER:
            with pdfplumber.open(pdf_file) as pdf:
                full_text = ""
                for page in pdf.pages:
                    t = page.extract_text() or ""
                    full_text += t + "\n"
                # Try to find ticket lines - common pattern: DOC NO, FARE, TAX, etc.
                # Example line: 057 1234567890 ...
                lines = full_text.split("\n")
                for line in lines:
                    # Look for ticket-like lines (10-13 digit numbers)
                    m = re.search(r'(\d{3})[-\s]?(\d{10,13})', line)
                    if m:
                        # Extract amounts - look for numbers with decimals
                        amounts = re.findall(r'\d+\.\d{2}', line)
                        tickets.append({
                            "raw": line.strip()[:200],
                            "doc_no": f"{m.group(1)}-{m.group(2)}",
                            "amounts": amounts
                        })
                summary_text = full_text[:5000]
        else:
            # Fallback text extraction via PyMuPDF if installed
            import fitz
            doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
            full_text = ""
            for page in doc:
                full_text += page.get_text()
            for line in full_text.split("\n"):
                m = re.search(r'(\d{3})[-\s]?(\d{10,13})', line)
                if m:
                    tickets.append({"raw": line.strip()[:200], "doc_no": f"{m.group(1)}-{m.group(2)}", "amounts": []})
            summary_text = full_text[:5000]
    except Exception as e:
        st.error(f"Parsing error: {e}")
    
    return tickets, summary_text

def create_excel(tickets, summary):
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "BSP-ISSUES"
    # Header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    headers = ["#", "Document Number", "Raw Line Extracted", "Amounts Found", "Remarks"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    # Data
    for i, t in enumerate(tickets, 1):
        ws.append([i, t.get("doc_no",""), t.get("raw",""), ", ".join(t.get("amounts",[])), ""])
    # Auto width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len+2, 60)
    # Second sheet - Summary
    ws2 = wb.create_sheet("SUMMARY")
    ws2.append(["IATA BSP Summary"])
    ws2.append([summary[:30000]])
    # Dev sheet
    ws3 = wb.create_sheet("INFO")
    ws3.append(["Developed by", "Mahmoud Amin"])
    ws3.append(["Tool", "IATA BSP Converter PRO"])
    ws3.append(["Version", "3.0 - Streamlit Edition"])
    ws3.append(["Tickets Found", len(tickets)])
    
    wb.save(output)
    output.seek(0)
    return output

# Main Layout - 2 columns fullscreen
col1, col2 = st.columns([2, 1], gap="large")

with col1:
    st.markdown('<div class="card">', unsafe_allow_html=True)
    st.markdown('<div class="step-label">STEP 1</div><h3 style="color:white; margin-top:5px;">📄 Select PDF File</h3>', unsafe_allow_html=True)
    uploaded = st.file_uploader("Drag & drop your EG_FCAGBILLDET PDF here", type=["pdf"], label_visibility="collapsed")
    
    st.markdown('<div style="height:20px"></div><div class="step-label">STEP 2</div><h3 style="color:white; margin-top:5px;">⚙️ Convert & Download</h3>', unsafe_allow_html=True)
    
    if uploaded:
        st.success(f"✓ Selected: {uploaded.name} ({uploaded.size//1024} KB)")
        with st.spinner("Parsing PDF..."):
            tickets, summary = parse_pdf_to_data(uploaded)
        
        if tickets:
            st.metric("Tickets Found", len(tickets))
            with st.expander(f"Preview {min(10, len(tickets))} tickets"):
                st.dataframe(pd.DataFrame(tickets[:10]), use_container_width=True)
            
            excel_file = create_excel(tickets, summary)
            st.download_button(
                "⚡ GENERATE & DOWNLOAD EXCEL NOW",
                data=excel_file,
                file_name=uploaded.name.replace(".pdf","").replace(".PDF","") + "_CONVERTED.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            st.balloons()
        else:
            st.warning("No tickets detected with current parser. The file will still be converted with raw text extraction.")
            excel_file = create_excel([{"doc_no":"N/A","raw":summary[:200],"amounts":[]}], summary)
            st.download_button("Download Raw Extraction", data=excel_file, file_name="BSP_Extraction.xlsx")
    else:
        st.info("👆 Upload your FCAGBILLDET PDF to start. Example: EG_FCAGBILLDET_*.PDF")
    
    st.markdown('</div>', unsafe_allow_html=True)

with col2:
    st.markdown('<div class="card">', unsafe_allow_html=True)
    st.markdown('<div class="step-label">FILE INFO</div><h4 style="color:white;">About this tool</h4>', unsafe_allow_html=True)
    st.markdown("""
    <div class="info-box">
    <b>Supported:</b><br>
    • EG_FCAGBILLDET_*.PDF<br>
    • Any BSP Billing PDF<br>
    • Arabic & English<br><br>
    <b>Output:</b><br>
    • Excel (.xlsx)<br>
    • BSP-ISSUES sheet<br>
    • SUMMARY sheet<br>
    • INFO with developer<br><br>
    <b>How to get public link:</b><br>
    1. Upload this file to GitHub<br>
    2. Go to share.streamlit.io<br>
    3. Deploy -> You get a link like:<br>
    https://yourapp.streamlit.app<br>
    Works on mobile too!
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("---")
    st.markdown("**Developed by**<br><h3 style='color:#3B82F6; margin:0;'>Mahmoud Amin</h3><p style='color:#94A3B8;'>IATA BSP Specialist<br>© 2026 All Rights Reserved</p>", unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

st.markdown('<div class="footer">© 2026 Mahmoud Amin - IATA BSP Converter PRO v3.0 - Streamlit Cloud Ready | Built with ❤️ for travel agencies</div>', unsafe_allow_html=True)
